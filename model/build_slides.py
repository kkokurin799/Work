"""Собирает 3 слайда с расчётом по значениям из Excel-модели (базовый сценарий).

Запуск после build_model.py: python3 model/build_slides.py
"""

import os

from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from build_model import ATOMS, OUT as XLSX, TEAMS

OUT = os.path.join(os.path.dirname(XLSX), "brokerage_ai_restructuring_slides.pptx")

NAVY = RGBColor(0x1F, 0x4E, 0x78)
LIGHT = RGBColor(0xDD, 0xEB, 0xF7)
ZEBRA = RGBColor(0xF5, 0xF8, 0xFC)
ACCENT = RGBColor(0xC0, 0x50, 0x00)
DARK = RGBColor(0x26, 0x26, 0x26)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

wb = load_workbook(XLSX, data_only=True)
calc, struct, mix = wb["Расчёт"], wb["Структура"], wb["Процентовка"]


def pct(v):
    return f"{round(v * 100):.0f}%"


def num(v):
    return f"{v:.1f}".replace(".", ",")


def text(slide, x, y, w, h, runs, size=14, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(runs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        parts = line if isinstance(line, list) else [(line, False, DARK)]
        for t, bold, color in parts:
            run = para.add_run()
            run.text = t
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return box


def panel(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def table(slide, x, y, w, col_w, rows, size=12, total_rows=(), row_h=0.34):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                   Inches(w), Inches(row_h * len(rows)))
    tbl = shape.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, v in enumerate(row):
            c = tbl.cell(i, j)
            c.text = str(v)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
            para = c.text_frame.paragraphs[0]
            para.font.size = Pt(size)
            para.alignment = PP_ALIGN.LEFT if j == 0 or (isinstance(v, str) and len(v) > 12) else PP_ALIGN.CENTER
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = NAVY
                para.font.bold = True
                para.font.color.rgb = WHITE
            elif i in total_rows:
                c.fill.fore_color.rgb = LIGHT
                para.font.bold = True
                para.font.color.rgb = DARK
            else:
                c.fill.fore_color.rgb = ZEBRA if i % 2 == 0 else WHITE
                para.font.color.rgb = DARK
    return tbl


def slide_base(prs, title, subtitle, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    panel(s, 0, 0, 13.333, 1.05, NAVY)
    text(s, 0.45, 0.14, 12.4, 0.55, [[(title, True, WHITE)]], size=24)
    text(s, 0.45, 0.62, 12.4, 0.4, [[(subtitle, False, RGBColor(0xDD, 0xEB, 0xF7))]], size=14)
    text(s, 12.2, 7.05, 0.9, 0.3, [[(f"{n}/3", False, RGBColor(0x80, 0x80, 0x80))]], size=11, align=PP_ALIGN.RIGHT)
    return s


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

# ------------------------------------------------------------------ Слайд 1
blocks = [
    ("Аналитика и согласования", range(0, 3),
     "Протоколы встреч, разбор нормативки, черновик постановки и критериев приёмки", "Переговоры с заказчиком"),
    ("Разработка и сборка", range(3, 7),
     "Код по постановке, автоматическое ревью, анализ влияния, пайплайны", "Выбор решения"),
    ("Тестирование и ПСИ", range(7, 12),
     "Тест-кейсы, тестовые данные, автотесты, локализация дефектов", "Сами испытания с заказчиком"),
]
rows = [["Блок типовой задачи", "Доля часов", "Что делает ИИ", "Что ИИ почти не ускоряет", "Экономия"]]
for name, idx, does, doesnt in blocks:
    w = sum(ATOMS[i][4] for i in idx)
    saved = sum(ATOMS[i][4] * ATOMS[i][5] for i in idx)
    rows.append([name, pct(w), does, doesnt, f"−{pct(saved / w)}"])
rows.append(["Типовая задача в целом", "100%", "Заказные доработки, законодательство, развитие", "",
             f"−{pct(calc['C4'].value)}"])
s = slide_base(prs, "Как ИИ сокращает трудозатраты: формула и атомы работы",
               "Экономию считаем по атомам работы в часах, а не одним процентом на вид работ", 1)
table(s, 0.45, 1.3, 12.43, [3.1, 1.1, 4.6, 2.43, 1.2], rows, size=13, total_rows=(4,), row_h=0.52)

panel(s, 0.45, 4.15, 6.1, 2.75, LIGHT)
text(s, 0.65, 4.25, 5.8, 2.6, [
    [("Формула", True, NAVY)],
    [("Экономия вида работ:  E = Σ w × a", True, DARK)],
    "w — доля атома в часах, a — насколько ИИ сокращает часы атома",
    [("Нужно людей:  Σ FTE × (1 − E) × (1 + o) − R", True, DARK)],
    "o = 5% — ревью результатов агентов и доводка пайплайна",
    "R — меньше переделок благодаря Intent-команде",
], size=14)

other = [["Другие виды работ", "Экономия", "Почему"],
         ["Сопровождение", f"−{pct(calc['C5'].value)}", "Классификация и ответы — да, диагностика — частично"],
         ["Внедрение", f"−{pct(calc['C6'].value)}", "Плюс половина работы уходит в Intent-команду"],
         ["Пресейлы", f"−{pct(calc['C7'].value)}", "Встречи и демо ИИ не ускоряет"]]
table(s, 6.8, 4.15, 6.08, [1.9, 1.1, 3.08], other, size=12, row_h=0.42)
text(s, 6.8, 5.95, 6.08, 1.0, [
    [("Почему не 40–50%: ", True, ACCENT),
     ("ускорив вдвое всю разработку, получим только 15% общего эффекта. "
      "Предел задают согласования и ПСИ.", False, DARK)],
], size=13)

# ------------------------------------------------------------------ Слайд 2
s = slide_base(prs, f"Численность: {struct['H12'].value:.0f} → {struct['H11'].value:.0f} человек "
                    f"(−{pct(-struct['H13'].value)})",
               "Шаг 1: сколько людей нужно при сегодняшнем объёме работ", 2)
rows = [["Вид работ", "Сейчас", "Экономия", "Нужно"]]
for i in range(10, 16):
    rows.append([calc[f"A{i}"].value, num(calc[f"B{i}"].value), f"−{pct(calc[f'D{i}'].value)}",
                 num(calc[f"E{i}"].value)])
rows.append(["Накладные ИИ − переделки", "", "",
             f"+{num(calc['E17'].value - calc['E18'].value)}"])
rows.append(["Исполнителей", num(calc["B16"].value), f"−{pct(calc['E20'].value)}", num(calc["E19"].value)])
table(s, 0.45, 1.3, 6.2, [3.0, 0.95, 1.2, 1.05], rows, size=12, total_rows=(8,), row_h=0.36)

short = ["Intent-команда", "Продуктовый учёт по всем рынкам", "Внутренний, бух. и доверительный учёт",
         "НДФЛ", "Кор-команда по фининструментам (FA и MSA)", "Микросервисная (рыночные данные и др.)"]
rows = [["Команда", "РК / А / Р / Т", "Чел."]]
for i, name in enumerate(short, start=4):
    t = TEAMS[i - 4]
    comp = "5 экспертов" if i == 4 else f"{t[2]} / {t[3]} / {t[4]} / {t[5]}"
    rows.append([name, comp, f"{struct[f'H{i}'].value:.0f}"])
rows.append(["Роли вне команд", "", f"{struct['H10'].value:.0f}"])
rows.append(["Итого", "", f"{struct['H11'].value:.0f}"])
table(s, 6.9, 1.3, 5.98, [3.98, 1.3, 0.7], rows, size=12, total_rows=(8,), row_h=0.36)

rows = [["Сверка по ролям", "Сейчас в штате", "Нужно", "В структуре"],
        ["Аналитика (включая Intent-команду)", f"{calc['D35'].value:.0f}", num(calc["C35"].value), f"{struct['B16'].value:.0f}"],
        ["Разработка", f"{calc['D36'].value:.0f}", num(calc["C36"].value), f"{struct['B17'].value:.0f}"],
        ["Тестирование и DevOps", f"{calc['D37'].value:.0f}", num(calc["C37"].value), f"{struct['B18'].value:.0f}"],
        ["Руководители на задачах (половина времени)", "", "", num(struct["B19"].value)],
        ["Итого исполнителей", f"{calc['D38'].value:.0f}", num(calc["C38"].value), num(struct["B20"].value)]]
table(s, 0.45, 4.85, 7.6, [3.7, 1.3, 1.2, 1.4], rows, size=11, total_rows=(5,), row_h=0.33)
text(s, 8.3, 4.95, 4.58, 2.0, [
    [("Структура закрывает потребность: ", True, ACCENT),
     (f"{num(struct['B20'].value)} исполнителя при потребности {num(calc['C38'].value)}.", False, DARK)],
    [("ИИ сильнее всего сжимает разработку. ", True, DARK),
     ("Нехватку тестирования закрывают продуктовые инженеры: избыток разработки уходит в автотесты с агентами.", False, DARK)],
], size=13)

# ------------------------------------------------------------------ Слайд 3
s = slide_base(prs, "Процентовка работ TO-BE: заказных доработок вдвое больше",
               "Шаг 2: при тех же 42 людях меняем структуру работ, а не численность", 3)
levers = [
    f"Вся высвобожденная мощность; объём ×{num(mix['B15'].value)} к сегодняшнему",
    "ИИ −34%; запас +10% объёма на рост нормативки",
    "ИИ −34%; +10% объёма",
    "Половина — в Intent-команду; тиражные скрипты настройки и миграции",
    "Агент первой линии; вдвое меньше обращений за счёт качества (Intent-команда, автотесты)",
    "Типовой пакет пресейла и демо-стенд",
]
rows = [["Вид работ", "AS-IS", "Как меняется", "Нужно, чел.", "TO-BE"]]
for i, lever in zip(range(5, 11), levers):
    rows.append([mix[f"A{i}"].value, pct(mix[f"C{i}"].value), lever, num(mix[f"K{i}"].value), pct(mix[f"L{i}"].value)])
rows.append(["Итого", "100%", "", num(mix["K11"].value), "100%"])
table(s, 0.45, 1.3, 12.43, [3.0, 0.9, 6.23, 1.2, 1.1], rows, size=13, total_rows=(7,), row_h=0.5)

panel(s, 0.45, 5.4, 6.1, 1.5, LIGHT)
text(s, 0.65, 5.47, 5.8, 1.4, [
    [("Логика", True, NAVY)],
    "Обязательные потоки (законодательство, развитие) сохраняют объём и дешевеют на треть. "
    "Сопровождение, внедрение и пресейлы сжимаются агентами и типовыми решениями. "
    "Остаток мощности — в заказные доработки, то есть в выручку.",
], size=12)
text(s, 6.8, 5.4, 6.08, 1.6, [
    [("Главное условие — вдвое меньше обращений в сопровождение. ", True, ACCENT),
     (f"Если только −30%: сопровождение {pct(mix['B19'].value)}, "
      f"заказные доработки {pct(mix['B20'].value)}. Численность при этом не меняется.", False, DARK)],
    [("База процентовки: ", True, DARK),
     (f"{num(mix['B13'].value)} чел. в производственных командах (исполнители и половина времени руководителей). "
      "Intent-команда — отдельная функция.", False, DARK)],
], size=12)

prs.save(OUT)
print("saved", OUT)
